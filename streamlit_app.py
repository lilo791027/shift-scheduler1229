import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from io import BytesIO
from datetime import datetime
import re
from itertools import cycle

# ==========================================
# 系統開關 (控制是否看起來像壞掉)
# ==========================================
IS_OPEN = True  # 改為 False 網頁會顯示「系統崩潰」

# Streamlit 頁面設定 (必須是第一個指令)
st.set_page_config(page_title="班表處理器", layout="wide")

if not IS_OPEN:
    # 這裡偽裝成系統發生嚴重 Bug 的樣子
    st.error("### Internal Server Error (500)")
    st.markdown("---")
    st.error("**Critical Failure**: Unexpected token in JSON at position 0 (line 125)")
    
    # 模擬一段可怕的報錯訊息
    fake_error = f"""
    Traceback (most recent call last):
      File "/app/main.py", line 482, in <module>
        engine.start_service()
      File "/app/utils/handler.py", line 12, in start_service
        connection = database.connect(db_config)
    ConnectionError: [Errno 110] Connection timed out: '192.168.1.104:3306'
    
    Current System Time: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
    """
    st.code(fake_error, language="python")
    st.info("系統維護人員已收到報錯通知，將於修復完成後重新開放。")
    st.stop()  # 阻斷下方所有程式碼執行

# --------------------
# 模組 1：解除合併儲存格並填入原值
# --------------------
def unmerge_and_fill(ws):
    for merged in list(ws.merged_cells.ranges):
        value = ws.cell(merged.min_row, merged.min_col).value
        ws.unmerge_cells(str(merged))
        for row in ws[merged.coord]:
            for cell in row:
                cell.value = value

# --------------------
# 模組 2：整理班表資料
# --------------------
def consolidate_selected_sheets(wb, sheet_names):
    all_data = []
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        unmerge_and_fill(ws)
        try:
            clinic_name = str(ws.cell(row=1, column=1).value).strip()[:4]
        except:
            clinic_name = "未知診所"

        max_row = ws.max_row
        max_col = ws.max_column
        
        for r in range(1, max_row + 1):
            for c in range(2, max_col + 1):
                cell_value = ws.cell(r, c).value
                if isinstance(cell_value, datetime):
                    date_val = cell_value
                    i = r + 3
                    while i <= max_row:
                        shift_type = str(ws.cell(i, c).value).strip()
                        if shift_type in ["", "None"] or isinstance(ws.cell(i, c).value, datetime):
                            break
                        
                        if shift_type in ["早", "午", "晚"]:
                            i += 1
                            while i <= max_row:
                                cell_v = ws.cell(i, c).value
                                if isinstance(cell_v, datetime):
                                    break
                                val = str(cell_v).strip()
                                if val in ["早", "午", "晚"]:
                                    break
                                if val and val not in ["None", "nan", "="]:
                                    all_data.append([
                                        clinic_name,
                                        date_val.strftime("%Y/%m/%d"),
                                        shift_type,
                                        val
                                    ])
                                i += 1
                            i -= 1
                        i += 1
    df = pd.DataFrame(all_data, columns=["診所", "日期", "班別", "姓名"])
    return df

# --------------------
# 模組 3：建立班別分析表
# --------------------
def create_shift_analysis(df_shift: pd.DataFrame, df_emp: pd.DataFrame, shift_map: dict) -> pd.DataFrame:
    df_shift = df_shift.copy()
    df_emp = df_emp.copy()
    
    df_shift.columns = [str(c).replace(" ", "").replace("　", "").strip() for c in df_shift.columns]
    df_emp.columns = [str(c).replace(" ", "").replace("　", "").strip() for c in df_emp.columns]
    
    def get_col_name(df, keywords):
        for col in df.columns:
            for kw in keywords:
                if kw in col: return col
        return None

    col_map = {
        "姓名": get_col_name(df_emp, ["姓名"]),
        "編號": get_col_name(df_emp, ["編號", "工號"]), 
        "職稱": get_col_name(df_emp, ["職稱", "職務", "職位"]),
        "部門": get_col_name(df_emp, ["部門", "單位"]),
        "分類": get_col_name(df_emp, ["分類", "類別"]),
        "特殊早班": get_col_name(df_emp, ["特殊早班", "特權"])
    }
    
    emp_dict = {}
    for _, row in df_emp.iterrows():
        name_col = col_map["姓名"]
        if not name_col: continue

        raw_name = str(row.get(name_col, "")).strip()
        clean_name_key = raw_name.replace(" ", "").replace("　", "")
        
        if clean_name_key and clean_name_key not in ["nan", "None"]:
            emp_dict[clean_name_key] = [
                str(row.get(col_map["編號"], "")).strip(),
                str(row.get(col_map["部門"], "")).strip(),
                str(row.get(col_map["職稱"], "")).strip(),
                str(row.get(col_map["分類"], "")).strip(),
                str(row.get(col_map["特殊早班"], "")).strip()
            ]

    shift_dict = {}
    for _, row in df_shift.iterrows():
        raw_name = str(row.get("姓名", "")).strip()
        clean_name_key = raw_name.replace(" ", "").replace("　", "")
        
        clinic = str(row.get("診所", "")).strip()
        date_val = row.get("日期", "")
        shift_type = str(row.get("班別", "")).strip()
        
        if not raw_name or pd.isna(date_val): continue
        
        key = f"{clean_name_key}|{date_val}|{clinic}|{raw_name}"
        if key not in shift_dict: shift_dict[key] = set()
        shift_dict[key].add(shift_type)

    data_out = []
    for key, shifts in shift_dict.items():
        clean_name_key, date_val, clinic, original_name = key.split("|")
        
        emp_info = emp_dict.get(clean_name_key, ["", "", "", "", ""])
        emp_id, emp_dept, emp_title, emp_category, emp_early_special = emp_info
        
        shift_parts = [s for s in ["早", "午", "晚"] if s in shifts]
        shift_type_for_code = "".join(sorted(shift_parts, key=lambda x: {"早": 1, "午": 2, "晚": 3}.get(x, 9)))
        
        class_code = get_class_code(emp_category, emp_early_special, clinic, shift_type_for_code, shift_map)
        
        data_out.append([clinic, emp_id, emp_dept, original_name, emp_title, date_val, shift_type_for_code, class_code])

    df_analysis = pd.DataFrame(
        data_out,
        columns=["診所", "員工編號", "所屬部門", "姓名", "職稱", "日期", "班別", "班別代碼"]
    )
    return df_analysis

def get_class_code(emp_category, emp_early_special, clinic_name, shift_type, shift_map):
    region = "立丞" if re.search(r"立丞", str(clinic_name), re.IGNORECASE) else "板土中京"
    is_early_special = str(emp_early_special).strip().lower() in ["是", "true", "1", "checked"]

    if is_early_special and "早" in shift_type:
        if shift_type == "早": return "【員工】純早班"
        elif shift_type == "早午": return f"【員工】{region}純早、午班"
        elif shift_type == "早晚": return f"【員工】{region}純早、晚班"
        elif shift_type == "早午晚": return f"【員工】{region}純早午晚班"
    
    if shift_type == "早":
        if "醫師" in emp_category: return "★醫師★早班"
        elif "主管" in emp_category: return "◇主管◇早班"
        elif "員工" in emp_category: return "【員工】早班"

    if shift_type == "早午晚":
        return f"{emp_category}{region}全天班"
    
    base = shift_map.get(shift_type, shift_type)
    if not str(base).strip().endswith("班"): base += "班"
    return str(emp_category) + str(region) + str(base)

# --------------------
# 模組 4：建立班別總表 (嚴格刪除無員編 + 自動填補)
# --------------------
def create_shift_summary(df_analysis: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    if df_analysis.empty:
        return pd.DataFrame(), pd.DataFrame()
        
    df_analysis = df_analysis.copy()
    df_analysis["日期"] = pd.to_datetime(df_analysis["日期"], errors="coerce")
    df_analysis = df_analysis.dropna(subset=["日期"])
    all_dates = sorted(df_analysis["日期"].dt.strftime("%Y-%m-%d").unique())

    # 建立總表字典
    summary_dict = {}
    info_map = {} 

    for _, row in df_analysis.iterrows():
        emp_id = str(row["員工編號"]).strip()
        emp_name = str(row["姓名"]).strip()
        emp_title = str(row["職稱"]).strip()
        
        shift_date = row["日期"].strftime("%Y-%m-%d")
        summary_dict.setdefault((emp_id, emp_name), {})[shift_date] = row["班別代碼"]
        info_map[(emp_id, emp_name)] = emp_title

    data_out = []
    
    # 垃圾關鍵字清單
    invalid_keywords = ["義診", "單診", "盤點", "電打", "公務", "測試"]

    for (emp_id, emp_name), shifts in summary_dict.items():
        title = info_map.get((emp_id, emp_name), "")
        
        has_valid_id = emp_id and emp_id.lower() not in ["nan", "none", ""]
        is_invalid_name = any(k in emp_name for k in invalid_keywords)
        
        if not has_valid_id or is_invalid_name:
            continue
            
        is_doctor_or_pt = ("醫師" in title) or ("兼職" in title)
        should_fill = not is_doctor_or_pt
        
        leave_cycle = cycle(["{sta}", "{res}"])
        
        row = [emp_id, emp_name]
        for d in all_dates:
            val = shifts.get(d, "")
            is_empty = (val is None) or (str(val).strip() in ["", "nan", "None"])
            
            if is_empty:
                if should_fill:
                    val = next(leave_cycle) # 自動填入
                else:
                    val = "" # 醫師/兼職 保持空白
            
            row.append(val)
        data_out.append(row)

    cols = ["員工編號", "員工姓名"] + all_dates
    return pd.DataFrame(data_out, columns=cols), pd.DataFrame()

# --------------------
# Streamlit 主程式
# --------------------
st.title("班表處理器")
st.success("✅ 規則：【無員編】者直接刪除；非醫師/兼職的員工，自動填補空班。")

col1, col2 = st.columns(2)
with col1:
    shift_file = st.file_uploader("1. 上傳班表 (Excel)", type=["xlsx", "xlsm"])
with col2:
    employee_file = st.file_uploader("2. 上傳員工資料 (Excel 或 CSV)", type=["xlsx", "xlsm", "csv"])

if shift_file and employee_file:
    try:
        wb_shift = load_workbook(shift_file, data_only=True)
    except Exception as e:
        st.error(f"班表讀取失敗: {e}")
        st.stop()
        
    df_emp_raw = None
    try:
        if employee_file.name.lower().endswith('.csv'):
            df_emp_raw = pd.read_csv(employee_file)
        else:
            wb_emp = load_workbook(employee_file, data_only=True)
            emp_sheet_name = st.selectbox("選擇員工資料工作表", wb_emp.sheetnames)
            ws = wb_emp[emp_sheet_name]
            data = list(ws.values)
            if data:
                cols = [str(c).strip() for c in data[0]]
                df_emp_raw = pd.DataFrame(data[1:], columns=cols)
    except Exception as e:
        st.error(f"員工資料讀取失敗: {e}")
        st.stop()

    sheets = [s for s in wb_shift.sheetnames if s not in ["彙整結果", "班別分析", "班別總表"]]
    selected_sheets = st.multiselect("選擇班表工作表", sheets)

    if st.button("🚀 開始處理", type="primary"):
        if not selected_sheets:
            st.warning("請選擇至少一個工作表")
        elif df_emp_raw is None or df_emp_raw.empty:
            st.error("員工資料是空的")
        else:
            with st.spinner("資料處理中..."):
                df_shift = consolidate_selected_sheets(wb_shift, selected_sheets)
                shift_map = {"早": "早", "午": "午", "晚": "晚"}
                
                df_analysis = create_shift_analysis(df_shift, df_emp_raw, shift_map)
                df_summary, _ = create_shift_summary(df_analysis)
            
            st.subheader("📊 班別總表 (處理完成)")
            st.dataframe(df_summary, use_container_width=True)

            with BytesIO() as output:
                with pd.ExcelWriter(output, engine="openpyxl") as writer:
                    df_summary.to_excel(writer, sheet_name="班別總表", index=False)
                st.download_button("📥 下載 Excel 結果", output.getvalue(), "班別總表_最終版.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ==========================================
# 頁尾署名區塊 (已加入)
# ==========================================
st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: gray; font-size: 14px;'>"
    "Developed by <b> 林雅婷 </b> | © 2025 班表處理系統"
    "</div>", 
    unsafe_allow_html=True
)
